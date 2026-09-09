"""通过公共预览结果验证工作簿内容与资源边界。"""

from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font

from yuxi.utils.filepreview import render_preview


def test_xlsx_preview_preserves_sheets_values_merges_and_formula_warning():
    """预览保留内容和已保存结果，不执行公式。"""
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet["A1"] = "<script>alert(1)</script>"
    sheet["A1"].font = Font(bold=True)
    sheet.merge_cells("A1:B1")
    sheet["A2"] = 0.25
    sheet["A2"].number_format = "0.00%"
    sheet["B2"] = "=1+1"
    sheet["C2"] = "=10+20"
    book.create_sheet("说明")["A1"] = "第二张表"
    book.create_sheet("隐藏").sheet_state = "hidden"
    buffer = BytesIO()
    book.save(buffer)
    # 独立缓存 oracle 为 99，刻意与公式算术结果 30 不同，防止预览重算。
    output = BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>10+20</f><v></v>", b"<f>10+20</f><v>99</v>")
            target.writestr(item, data)
    result = render_preview("工作簿.XLSX", output.getvalue())
    assert result.supported is True
    assert result.preview_type == "spreadsheet"
    sheets = result.payload()["content"]["sheets"]
    assert [s["name"] for s in sheets] == ["销售", "说明"]
    assert sheets[0]["rows"][0][0]["text"] == "<script>alert(1)</script>"
    assert sheets[0]["rows"][0][0]["colspan"] == 2
    assert sheets[0]["rows"][0][0]["style"]["fontWeight"] == "bold"
    assert sheets[0]["rows"][1][0]["text"] == "25.00%"
    assert "未保存" in sheets[0]["rows"][1][1]["text"]
    assert sheets[0]["rows"][1][2]["text"] == "99"
    assert sheets[1]["rows"][0][0]["text"] == "第二张表"


def test_xls_preview_preserves_legacy_workbook():
    """真实 BIFF 工作簿保留工作表、格式与合并关系。"""
    from pathlib import Path

    content = (Path(__file__).parents[2] / "fixtures/office/readonly.xls").read_bytes()
    result = render_preview("旧版.xls", content)
    assert result.supported is True
    sheets = result.content["sheets"]
    assert [s["name"] for s in sheets] == ["销售", "说明"]
    assert sheets[0]["rows"][0][0]["text"] == "销售汇总"
    assert sheets[0]["rows"][0][0]["colspan"] == 2
    assert sheets[0]["rows"][1][0]["text"] == "25.00%"
    assert "未保存" in sheets[0]["rows"][1][1]["text"]


def test_spreadsheet_rejects_corruption_and_oversized_grid_without_truncation():
    """损坏文件和稀疏超大范围均明确失败，不返回部分成功。"""
    for name in ["broken.xlsx", "encrypted.xls"]:
        result = render_preview(name, b"not an office document")
        assert result.supported is False
        assert "下载" in result.message
    book = Workbook()
    book.active["A100001"] = "不能静默遗漏"
    buffer = BytesIO()
    book.save(buffer)
    result = render_preview("large.xlsx", buffer.getvalue())
    assert result.supported is False
    assert result.truncated is False
    assert "100000" in result.message


def test_spreadsheet_rejects_zip_expansion_and_oversized_input():
    """压缩包展开预算和原文件大小分别受限。"""
    from zipfile import ZIP_DEFLATED
    from yuxi.utils.filepreview import MAX_BINARY_PREVIEW_SIZE_BYTES

    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (64 * 1024 * 1024 + 1))
    result = render_preview("bomb.xlsx", buffer.getvalue())
    assert result.supported is False
    assert "64 MB" in result.message
    result = render_preview("large.xls", b"x" * (MAX_BINARY_PREVIEW_SIZE_BYTES + 1))
    assert result.supported is False
    assert "30 MB" in result.message


def test_spreadsheet_timeout_stops_parser(monkeypatch):
    """超时杀死真实解析子进程并返回可恢复提示。"""
    from yuxi.utils import spreadsheet_preview

    monkeypatch.setattr(spreadsheet_preview, "SPREADSHEET_TIMEOUT_SECONDS", 0.000001)
    result = render_preview("slow.xlsx", b"content")
    assert result.supported is False
    assert "超时" in result.message


def test_spreadsheet_rejects_excessive_zip_entries():
    """小体积压缩包也必须遵守文件项数量限制。"""
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for index in range(10001):
            archive.writestr(f"entry-{index}", b"")
    result = render_preview("many-entries.xlsx", buffer.getvalue())
    assert result.supported is False
    assert result.truncated is False
    assert "文件项过多" in result.message


def test_spreadsheet_rejects_oversized_response_without_partial_content():
    """未超过输入和网格预算的长文本仍须遵守响应体上限。"""
    book = Workbook()
    for row in range(1, 601):
        book.active.cell(row, 1, "x" * 30000)
    buffer = BytesIO()
    book.save(buffer)
    result = render_preview("long-text.xlsx", buffer.getvalue())
    assert result.supported is False
    assert result.content is None
    assert result.truncated is False
    assert "预览内容过大" in result.message


def test_xls_preview_uses_saved_formula_result_without_recalculation():
    """保存值故意与 1+1 的算术结果不同，证明旧格式也不重算。"""
    import struct
    from pathlib import Path

    content = (Path(__file__).parents[2] / "fixtures/office/readonly.xls").read_bytes()
    empty_cache = b"\x03\x00\x00\x00\x00\x00\xff\xff"
    assert content.count(empty_cache) == 1
    content = content.replace(empty_cache, struct.pack("<d", 99.0))
    result = render_preview("cached.xls", content)
    assert result.supported is True
    assert result.content["sheets"][0]["rows"][1][1]["text"] == "99"


def test_spreadsheet_rejects_merges_exceeding_native_table_spans():
    """浏览器会缩短超大合并跨度，预览必须明确拒绝而非错列。"""
    for area in ["A1:ALM1", "A1:A65535"]:
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "合并区域"
        sheet.merge_cells(area)
        buffer = BytesIO()
        book.save(buffer)
        result = render_preview("wide-merge.xlsx", buffer.getvalue())
        assert result.supported is False
        assert result.content is None
        assert "合并跨度" in result.message
