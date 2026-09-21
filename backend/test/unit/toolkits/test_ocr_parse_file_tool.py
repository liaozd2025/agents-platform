from __future__ import annotations

from types import SimpleNamespace
from pathlib import Path
from io import BytesIO

import pytest
from docx import Document
from langchain_core.messages import ToolMessage

from yuxi.agents.backends import create_agent_filesystem_middleware
from yuxi.agents.buildin.subagent.graph import _SubAgentToolFilterMiddleware
from yuxi.agents.middlewares.model_input import ImageInputCompatibilityMiddleware
from yuxi.services import ocr_service
from yuxi.agents.toolkits.buildin import tools as buildin_tools
from yuxi.agents.toolkits.buildin.tools import ocr_parse_file

pytestmark = pytest.mark.unit


def _patch_sandbox_backend(monkeypatch: pytest.MonkeyPatch, files: dict[str, bytes]):
    class FakeBackend:
        def __init__(self, **kwargs):
            assert kwargs["create_if_missing"] is True
            self.scope = kwargs

        def download_authorized_file_to_path(self, path, target_path, max_bytes):
            if path not in files:
                raise ValueError("not a regular file")
            content = files[path]
            assert len(content) <= max_bytes
            Path(target_path).write_bytes(content)
            return len(content)

        def regular_file_exists(self, path):
            return path in files

        def upload_authorized_file_from_path(self, path, source_path):
            files[path] = Path(source_path).read_bytes()

    monkeypatch.setattr(buildin_tools, "ProvisionerSandboxBackend", FakeBackend, raising=False)
    return files


def _runtime(
    *,
    thread_id: str = "thread-1",
    uid: str = "user-1",
) -> SimpleNamespace:
    configurable = {
        "thread_id": thread_id,
        "runtime_scope_id": thread_id,
        "workdir_relative_path": "projects/11111111-1111-4111-8111-111111111111",
        "workdir_path": "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111",
        "uid": uid,
    }
    return SimpleNamespace(
        config={"configurable": configurable},
        context=SimpleNamespace(
            thread_id=thread_id,
            runtime_scope_id=thread_id,
            workdir_relative_path="projects/11111111-1111-4111-8111-111111111111",
            workdir_path="/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111",
            uid=uid,
        ),
        state={},
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("subagent", [False, True])
@pytest.mark.parametrize("extension", ["doc", "docx"])
async def test_office_tool_passes_execution_filters_and_parses_word(monkeypatch, subagent, extension):
    """主、子 Agent 经真实执行过滤后解析 Word，并保存完整 Markdown。"""
    _mock_system_options(monkeypatch)
    source = BytesIO()
    document = Document()
    document.add_paragraph("Office backend content")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Quantity"
    table.cell(1, 0).text = "Sample"
    table.cell(1, 1).text = "37"
    document.save(source)
    source_bytes = source.getvalue()
    expected = ("Office backend content", "37")
    if extension == "doc":
        source_bytes = (Path(__file__).parents[2] / "data/legacy_word.doc").read_bytes()
        expected = ("中文正文 DOC-ROUNDTRIP-2026", "42")
    path = f"{_runtime().context.workdir_path}/uploads/report.{extension}"
    files = _patch_sandbox_backend(monkeypatch, {path: source_bytes})
    tools = {tool.name: tool for tool in ImageInputCompatibilityMiddleware().tools}
    assert "ocr_parse_file" in tools
    assert "旧 DOC 由后端自动转换后解析，可直接传入原件" in tools["ocr_parse_file"].description
    filesystem = create_agent_filesystem_middleware(backend=SimpleNamespace())
    request = SimpleNamespace(tool_call={"name": "ocr_parse_file", "id": "parse-office", "args": {"file_path": path}})

    async def parse(request):
        result = await tools[request.tool_call["name"]].coroutine(file_path=path, runtime=_runtime())
        assert expected[0] in result["preview"]
        assert expected[1] in files[result["parsed_path"]].decode()
        assert files[path] == source_bytes
        return ToolMessage(content=result["parsed_path"], tool_call_id="parse-office")

    async def execute(request):
        return await filesystem.awrap_tool_call(request, parse)

    result = (
        await _SubAgentToolFilterMiddleware().awrap_tool_call(request, execute) if subagent else await execute(request)
    )
    assert result.status == "success"
    assert result.content.endswith("/outputs/ocr/report.md")


@pytest.mark.asyncio
async def test_ocr_parse_file_writes_markdown_to_outputs(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("YUXI_USER_DATA_DIR", str(tmp_path / "threads"))
    _mock_system_options(monkeypatch)

    def resolve_engine(engine_id, default_engine):
        del default_engine
        return engine_id

    monkeypatch.setattr(ocr_service, "resolve_ocr_engine_id", resolve_engine)
    thread_id = "thread-1"
    uid = "user-1"
    source_virtual_path = "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111/scan.png"
    sandbox_files = _patch_sandbox_backend(monkeypatch, {source_virtual_path: b"fake image"})
    captured: dict[str, object] = {}

    async def fake_parse_document(source: str, params: dict | None = None, db=None) -> str:
        del db
        captured["source"] = source
        captured["params"] = params
        return "识别结果\n" + ("长文本" * 500)

    monkeypatch.setattr(ocr_service, "parse_document", fake_parse_document)

    result = await ocr_parse_file.coroutine(
        file_path=source_virtual_path,
        ocr_engine="mineru_ocr",
        runtime=_runtime(thread_id=thread_id, uid=uid),
    )

    output_virtual_path = "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111/outputs/ocr/scan.md"
    assert sandbox_files[output_virtual_path].decode("utf-8").startswith("识别结果")
    assert result["source_path"] == source_virtual_path
    assert result["parsed_path"] == output_virtual_path
    assert result["ocr_engine"] == "mineru_ocr"
    assert result["char_count"] == len(sandbox_files[output_virtual_path].decode("utf-8"))
    assert result["truncated"] is True
    assert len(result["preview"]) <= 1200
    assert Path(str(captured["source"])).suffix == ".png"
    assert captured["params"] == {"ocr_engine": "mineru_ocr"}


@pytest.mark.asyncio
async def test_ocr_parse_file_uses_default_engine(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("YUXI_USER_DATA_DIR", str(tmp_path / "threads"))
    _mock_system_options(monkeypatch)

    def resolve_engine(engine_id, default_engine):
        assert engine_id is None
        assert default_engine == "rapid_ocr"
        return "rapid_ocr"

    monkeypatch.setattr(ocr_service, "resolve_ocr_engine_id", resolve_engine)
    thread_id = "thread-1"
    uid = "user-1"
    source_virtual_path = "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111/uploads/upload.pdf"
    _patch_sandbox_backend(monkeypatch, {source_virtual_path: b"fake pdf"})
    captured: dict[str, object] = {}

    async def fake_parse_document(source: str, params: dict | None = None, db=None) -> str:
        del source, db
        captured["params"] = params
        return "OCR content"

    monkeypatch.setattr(ocr_service, "parse_document", fake_parse_document)

    result = await ocr_parse_file.coroutine(
        file_path=source_virtual_path,
        runtime=_runtime(thread_id=thread_id, uid=uid),
    )

    assert result["ocr_engine"] == "rapid_ocr"
    assert captured["params"] == {"ocr_engine": "rapid_ocr"}


@pytest.mark.asyncio
async def test_ocr_parse_file_accepts_disable_for_pdf(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("YUXI_USER_DATA_DIR", str(tmp_path / "threads"))
    _mock_system_options(monkeypatch)
    thread_id = "thread-1"
    uid = "user-1"
    source_virtual_path = "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111/uploads/text-layer.pdf"
    _patch_sandbox_backend(monkeypatch, {source_virtual_path: b"fake pdf"})
    captured: dict[str, object] = {}

    async def fake_parse_document(source: str, params: dict | None = None, db=None) -> str:
        del source, db
        captured["params"] = params
        return "PDF text layer"

    monkeypatch.setattr(ocr_service, "parse_document", fake_parse_document)

    result = await ocr_parse_file.coroutine(
        file_path=source_virtual_path,
        ocr_engine="disable",
        runtime=_runtime(thread_id=thread_id, uid=uid),
    )

    assert result["ocr_engine"] == "disable"
    assert captured["params"] == {"ocr_engine": "disable"}


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "file_path",
    [
        "/etc/passwd",
        "/home/gem/user-data/../secrets.png",
    ],
)
async def test_ocr_parse_file_rejects_path_outside_user_data(
    tmp_path, monkeypatch: pytest.MonkeyPatch, file_path: str
) -> None:
    monkeypatch.setenv("YUXI_USER_DATA_DIR", str(tmp_path / "threads"))
    _mock_system_options(monkeypatch)

    with pytest.raises(ValueError, match="只允许解析"):
        await ocr_parse_file.coroutine(file_path=file_path, runtime=_runtime())


@pytest.mark.asyncio
async def test_ocr_parse_file_rejects_directory(tmp_path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("YUXI_USER_DATA_DIR", str(tmp_path / "threads"))
    _mock_system_options(monkeypatch)
    thread_id = "thread-1"
    uid = "user-1"
    dir_virtual_path = "/home/gem/user-data/projects/11111111-1111-4111-8111-111111111111/directory"
    _patch_sandbox_backend(monkeypatch, {})

    with pytest.raises(ValueError, match="不存在或不是普通文件"):
        await ocr_parse_file.coroutine(file_path=dir_virtual_path, runtime=_runtime(thread_id=thread_id, uid=uid))


def _mock_system_options(monkeypatch: pytest.MonkeyPatch) -> None:
    from yuxi.config.options import Option, system_options

    async def get_options(option, _db=None):
        assert option is system_options
        return {"default_ocr_engine": "rapid_ocr"}

    monkeypatch.setattr(Option, "get", get_options)
