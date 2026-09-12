"""在有资源上限的子进程内读取 Excel 保存状态，不计算公式。"""

from __future__ import annotations

import io
import json
import math
import re
import subprocess
import sys
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile

MAX_SPREADSHEET_CELLS = 100_000
MAX_SPREADSHEET_EXPANDED_BYTES = 64 * 1024 * 1024
SPREADSHEET_TIMEOUT_SECONDS = 30


class SpreadsheetLimitError(ValueError):
    """工作簿超出安全预览范围。"""


def preview_spreadsheet(suffix: str, content: bytes) -> dict:
    """隔离不可信解析，并返回可直接放入预览响应的结果。"""
    try:
        result = subprocess.run(
            [sys.executable, str(Path(__file__).resolve()), suffix],
            input=content,
            capture_output=True,
            timeout=SPREADSHEET_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return {"error": "Excel 预览超时，请下载原文件查看"}
    if result.returncode != 0:
        return {"error": "Excel 解析失败或超出资源限制，请下载原文件查看"}
    return json.loads(result.stdout)


def _cell_text(value, number_format: str) -> str:
    """显示常见数字格式，其他格式保留原值供只读查看。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if not isinstance(value, (int, float)):
        return str(value)
    if not math.isfinite(value):
        return str(value)
    # ponytail: 仅格式化常用十进制、千分位、百分数与货币；复杂自定义格式保留原值和格式提示。
    match = re.fullmatch(r"([¥￥$€£]?)(#,##0|0)(?:\.(0{1,10}))?(%)?", number_format)
    if match:
        currency, integer, decimals, percent = match.groups()
        return (
            currency
            + format(value * (100 if percent else 1), f"{',' if ',' in integer else ''}.{len(decimals or '')}f")
            + (percent or "")
        )
    return str(int(value)) if float(value).is_integer() else str(value)


def _finish_sheet(name: str, rows: list, merges: list) -> dict:
    """把合并区域归并到左上单元格，保留表格坐标。"""
    for r1, r2, c1, c2 in merges:
        if r2 - r1 > 65534 or c2 - c1 > 1000:
            raise SpreadsheetLimitError("Excel 合并跨度超过浏览器预览限制，请下载原文件查看")
        if r2 > len(rows) or c2 > (len(rows[0]) if rows else 0):
            raise SpreadsheetLimitError("Excel 合并区域超出工作表范围")
        rows[r1][c1]["rowspan"] = r2 - r1
        rows[r1][c1]["colspan"] = c2 - c1
        for row in range(r1, r2):
            for col in range(c1, c2):
                if (row, col) != (r1, c1):
                    rows[row][col] = None
    return {"name": name, "rows": rows}


def _check_dimensions(total: int, rows: int, cols: int) -> int:
    """拒绝超过浏览器展示预算的工作簿，避免静默截断。"""
    total += rows * cols
    if total > MAX_SPREADSHEET_CELLS:
        raise SpreadsheetLimitError("Excel 展开区域超过 100000 个单元格，请下载原文件查看")
    return total


def _xlsx_color(color, palette: list[str]) -> str | None:
    """将工作簿调色板和主题颜色转成受限 RGB 值。"""
    import colorsys
    from openpyxl.styles.colors import COLOR_INDEX

    if color is None:
        return None
    if color.type == "rgb":
        rgb = color.rgb[-6:]
    elif color.type == "indexed" and 0 <= color.indexed < len(COLOR_INDEX):
        rgb = COLOR_INDEX[color.indexed][-6:]
    elif color.type == "theme" and 0 <= color.theme < len(palette):
        rgb = palette[color.theme]
    else:
        return None
    if not re.fullmatch(r"[0-9a-fA-F]{6}", rgb):
        return None
    if color.tint:
        hue, light, saturation = colorsys.rgb_to_hls(*(int(rgb[i : i + 2], 16) / 255 for i in (0, 2, 4)))
        light = light * (1 + color.tint) if color.tint < 0 else light * (1 - color.tint) + color.tint
        rgb = "".join(f"{round(value * 255):02x}" for value in colorsys.hls_to_rgb(hue, light, saturation))
    return "#" + rgb


def _read_xlsx(content: bytes) -> list:
    """读取可见工作表的缓存值及基本样式。"""
    import openpyxl
    from defusedxml.ElementTree import fromstring

    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(item.file_size for item in entries) > MAX_SPREADSHEET_EXPANDED_BYTES:
            raise SpreadsheetLimitError("Excel 解压后超过 64 MB 或文件项过多，请下载原文件查看")
    formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    values = openpyxl.load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    palette = []
    if formulas.loaded_theme:
        theme = fromstring(formulas.loaded_theme)
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        scheme = theme.find(f"{ns}themeElements/{ns}clrScheme")
        for name in [
            "lt1",
            "dk1",
            "lt2",
            "dk2",
            "accent1",
            "accent2",
            "accent3",
            "accent4",
            "accent5",
            "accent6",
            "hlink",
            "folHlink",
        ]:
            entry = scheme.find(ns + name)[0]
            palette.append(entry.get("lastClr") or entry.get("val") or "")
    sheets = []
    total = 0
    try:
        for sheet in formulas.worksheets:
            if sheet.sheet_state != "visible":
                continue
            total = _check_dimensions(total, sheet.max_row, sheet.max_column)
            rows = []
            for row in sheet.iter_rows():
                cells = []
                for cell in row:
                    value = values[sheet.title].cell(cell.row, cell.column).value
                    text = _cell_text(value, cell.number_format)
                    if cell.data_type == "f" and value is None:
                        text = "公式结果为空或未保存"
                    style = {
                        "fontWeight": "bold" if cell.font.bold else "normal",
                        "fontStyle": "italic" if cell.font.italic else "normal",
                        "textAlign": cell.alignment.horizontal
                        if cell.alignment.horizontal in {"left", "center", "right", "justify"}
                        else "left",
                    }
                    for key, color in (
                        ("color", cell.font.color),
                        ("backgroundColor", cell.fill.fgColor if cell.fill.patternType == "solid" else None),
                    ):
                        rgb = _xlsx_color(color, palette)
                        if rgb:
                            style[key] = rgb
                    cells.append({"text": text, "style": style, "format": cell.number_format})
                rows.append(cells)
            merges = [(r.min_row - 1, r.max_row, r.min_col - 1, r.max_col) for r in sheet.merged_cells.ranges]
            sheets.append(_finish_sheet(sheet.title, rows, merges))
    finally:
        formulas.close()
        values.close()
    return sheets


def _xls_empty_formulas(content: bytes, biff_version: int) -> set:
    """识别 BIFF 空公式缓存；xlrd 的 Cell 会将其与普通空文本合并。"""
    import struct
    from xlrd.compdoc import CompDoc

    stream = content
    if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        document = CompDoc(content, logfile=io.StringIO())
        stream = document.get_named_stream("Workbook") or document.get_named_stream("Book")
    empty = set()
    offset, sheet_index = 0, -1
    while offset + 4 <= len(stream):
        code, size = struct.unpack_from("<HH", stream, offset)
        data = stream[offset + 4 : offset + 4 + size]
        offset += size + 4
        if code in {0x0809, 0x0409, 0x0209, 0x0009} and len(data) >= 4:
            if struct.unpack_from("<H", data, 2)[0] == 0x0010:
                sheet_index += 1
        elif code in {0x0006, 0x0206, 0x0406} and len(data) >= 15:
            result_start = 7 if biff_version < 30 else 6
            cached = data[result_start : result_start + 8]
            if cached[6:8] == b"\xff\xff" and cached[0] == 3:
                row, col = struct.unpack_from("<HH", data)
                empty.add((sheet_index, row, col))
    return empty


def _read_xls(content: bytes) -> list:
    """直接读取 BIFF 缓存值和基本格式，不经过 Office 转换。"""
    import xlrd

    book = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True, logfile=io.StringIO())
    sheets, total = [], 0
    try:
        empty_formulas = _xls_empty_formulas(content, book.biff_version)
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            if sheet.visibility:
                book.unload_sheet(index)
                continue
            row_count = max([sheet.nrows] + [r2 for r1, r2, c1, c2 in sheet.merged_cells])
            col_count = max([sheet.ncols] + [c2 for r1, r2, c1, c2 in sheet.merged_cells])
            total = _check_dimensions(total, row_count, col_count)
            rows = []
            for row in range(row_count):
                cells = []
                for col in range(col_count):
                    cell = sheet.cell(row, col) if row < sheet.nrows and col < sheet.ncols else xlrd.empty_cell
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#ERROR!")
                    xf = book.xf_list[cell.xf_index or 0]
                    font = book.font_list[xf.font_index]
                    number_format = book.format_map[xf.format_key].format_str
                    text = _cell_text(value, number_format)
                    if (index, row, col) in empty_formulas:
                        text = "公式结果为空或未保存"
                    style = {
                        "fontWeight": "bold" if font.bold else "normal",
                        "fontStyle": "italic" if font.italic else "normal",
                        "textAlign": {1: "left", 2: "center", 3: "right", 5: "justify"}.get(
                            xf.alignment.hor_align, "left"
                        ),
                    }
                    for key, color_index in (
                        ("color", font.colour_index),
                        (
                            "backgroundColor",
                            xf.background.pattern_colour_index if xf.background.fill_pattern == 1 else None,
                        ),
                    ):
                        rgb = book.colour_map.get(color_index)
                        if rgb:
                            style[key] = f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"
                    cells.append({"text": text, "style": style, "format": number_format})
                rows.append(cells)
            sheets.append(_finish_sheet(sheet.name, rows, sheet.merged_cells))
            book.unload_sheet(index)
    finally:
        book.release_resources()
    return sheets


def _main() -> None:
    """进程边界限制内存、CPU、输入和输出，失败仅返回固定错误。"""
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (1024 * 1024 * 1024, 1024 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    try:
        content = sys.stdin.buffer.read(30 * 1024 * 1024 + 1)
        if len(content) > 30 * 1024 * 1024:
            raise SpreadsheetLimitError("文件过大，当前仅支持 30 MB 以内的文件预览")
        sheets = _read_xls(content) if sys.argv[1] == ".xls" else _read_xlsx(content)
        payload = json.dumps({"sheets": sheets}, ensure_ascii=False)
        if len(payload.encode("utf-8")) > 16 * 1024 * 1024:
            raise SpreadsheetLimitError("Excel 预览内容过大，请下载原文件查看")
    except SpreadsheetLimitError as exc:
        payload = json.dumps({"error": str(exc)}, ensure_ascii=False)
    except Exception:
        payload = json.dumps({"error": "Excel 文件损坏、已加密或无法解析，请下载原文件查看"}, ensure_ascii=False)
    sys.stdout.write(payload)


if __name__ == "__main__":
    _main()
